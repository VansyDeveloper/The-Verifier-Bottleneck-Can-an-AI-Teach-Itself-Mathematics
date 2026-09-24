"""Build a concise university-facing report from audited September results."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "verifier_bottleneck_work_report_2026_09_22_24_ru.docx"
FIGURE = ROOT / "reports" / "trajectory_diversity_20260922" / "figures" / "selection_hitk_1_125.png"


def font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph.add_run()._r.addnext(field)


def add_body(doc: Document, text: str):
    paragraph = doc.add_paragraph(text, style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.widow_control = True
    return paragraph


def add_heading(doc: Document, text: str):
    paragraph = doc.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_caption(doc: Document, text: str, before: bool = False):
    paragraph = doc.add_paragraph(text, style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = before
    return paragraph


def format_cell(cell, text: str, header: bool, alternate: bool, align) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        line = borders.find(tag)
        if line is None:
            line = OxmlElement(f"w:{edge}")
            borders.append(line)
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), "D9D9D9")
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, amount in (("top", 85), ("bottom", 85), ("left", 105), ("right", 105)):
        tag = qn(f"w:{edge}")
        node = margins.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
    if header or alternate:
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E9EDF2" if header else "F7F8FA")
        tc_pr.append(shade)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.bold = header


def add_status_table(doc: Document) -> None:
    add_caption(doc, "Таблица 1. Выполнение эксперимента с исключёнными парами операций", before=True)
    headers = ("Число исключённых пар", "Допустимых наборов", "Завершённых наборов", "Принятых обучений")
    rows = (
        ("1", "5", "5", "30"),
        ("2", "5", "5", "30"),
        ("3", "5", "3", "18"),
        ("5", "4", "0", "0"),
        ("Всего", "19", "13", "78"),
    )
    widths = (Cm(4.1), Cm(4.0), Cm(4.1), Cm(4.3))
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = width
    for row_index, values in enumerate((headers, *rows)):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for i, value in enumerate(values):
            row.cells[i].width = widths[i]
            format_cell(
                row.cells[i],
                value,
                header=row_index == 0,
                alternate=row_index % 2 == 0,
                align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER,
            )
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main() -> None:
    if not FIGURE.is_file():
        raise FileNotFoundError(FIGURE)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.footer_distance = Cm(1.2)
    add_page_number(section)

    styles = doc.styles
    font(styles["Normal"], 14)
    styles["Normal"].paragraph_format.first_line_indent = Cm(1.25)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.5
    font(styles["Title"], 16, bold=True)
    styles["Title"].paragraph_format.space_after = Pt(9)
    title_style_properties = styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)
    font(styles["Heading 1"], 14, bold=True)
    styles["Heading 1"].paragraph_format.space_before = Pt(13)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 1"].paragraph_format.first_line_indent = Cm(0)
    font(styles["Caption"], 11)
    styles["Caption"].paragraph_format.first_line_indent = Cm(0)
    styles["Caption"].paragraph_format.line_spacing = 1.0
    styles["Caption"].paragraph_format.space_before = Pt(6)
    styles["Caption"].paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Отчёт о результатах исследования композиционного обобщения программ")
    title_border = OxmlElement("w:pBdr")
    title_bottom = OxmlElement("w:bottom")
    title_bottom.set(qn("w:val"), "nil")
    title_border.append(title_bottom)
    title._p.get_or_add_pPr().append(title_border)
    period = doc.add_paragraph("С 22 по 24 сентября 2026 года")
    period.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period.paragraph_format.first_line_indent = Cm(0)
    period.paragraph_format.space_after = Pt(12)
    for run in period.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    add_body(doc, "Новая экспериментальная серия началась 22 сентября в 20:06 МСК. За отчётный период завершены парное сравнение способов отбора обучающих траекторий, 13 наборов эксперимента с исключёнными парами операций и парное сравнение обучения с бонусом энтропии. Эти три опыта дали 96 принятых обучений и 96 полных оценок. Ранее сохранённые ранжирования программ и журналы GRPO анализировались отдельно и не включены в число новых обучений.")

    add_heading(doc, "Анализ сохранённых ранжирований и журналов GRPO")
    add_body(doc, "В шести ранее проведённых повторах проверены оценки всех 125 допустимых программ для каждой из 1000 задач глубины 3. Средний Hit@32 составил 0,4532 для атомарного контроля и 0,7272 после композиционного обучения. Условная энтропия распределения по 125 программам равна 0,765 и 3,892 нат соответственно. Эти значения описывают конечные модели. Они не измеряют энтропию полного словаря модели и не показывают её изменение во время обучения.")
    add_body(doc, "Отдельно пересчитаны шесть ранее сохранённых журналов GRPO по 400 шагов. В 2400 группах получено 19 200 ответов. При выборке IID 1121 групп содержали только неверные ответы, 38 были смешанными и 41 содержала только верные ответы. При выборке prefix-balanced соответствующие числа равны 896, 304 и 0. Средняя выборочная энтропия восьми ответов равна 0,298 и 1,493 нат. По блокам из 25 шагов устойчивого монотонного снижения разнообразия не обнаружено. Выборочную энтропию ответов нельзя приравнивать к энтропии политики.")

    add_heading(doc, "Отбор обучающих траекторий")
    add_body(doc, "В новом парном опыте шесть пар моделей обучались на одном исходном атомарном checkpoint. Каждая ветвь получала 750 композиционных траекторий. Случайный и структурированный наборы были выровнены по глубине задачи, параметру p и числу обучающих токенов. Структурированный отбор повысил заранее заданный показатель разнообразия обучающего набора во всех шести парах. После обучения обе ветви оценены на одних и тех же 1000 новых задачах с точным ранжированием 125 программ.")
    add_body(doc, "Средний Hit@32 равен 0,6155 при случайном отборе и 0,6130 при структурированном. Средняя парная разность составляет -0,0025. Её 95-процентный t-интервал лежит от -0,0227 до 0,0177, точное двустороннее значение p равно 0,8125. При данном бюджете структурированный отбор не показал преимущества по основному показателю. Отсутствие подтверждённого преимущества не означает равенства двух способов отбора.")
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.first_line_indent = Cm(0)
    figure.paragraph_format.space_before = Pt(5)
    figure.paragraph_format.space_after = Pt(0)
    figure.paragraph_format.keep_with_next = True
    shape = figure.add_run().add_picture(str(FIGURE), width=Cm(15.4))
    shape._inline.docPr.set("descr", "Две кривые Hit@K для случайного и структурированного отбора по шести обучениям, K от 1 до 125")
    add_caption(doc, "Рисунок 1. Доля решённых задач среди первых K программ, среднее по шести обучениям")

    add_heading(doc, "Исключение пар операций")
    add_body(doc, "Сравниваются обучение без заданных соседних пар операций и контроль, из которого удалено столько же случайных траекторий. Объём выборки и распределения по глубине и параметру p выровнены. Каждый набор включает три парных обучающих seed, то есть шесть обучений и шесть полных оценок. Основной показатель есть разность Hit@32 на задачах с исключёнными парами.")
    add_status_table(doc)
    add_body(doc, "Из 20 заранее намеченных наборов допустимы 19. Набор с пятью исключёнными парами под номером 5 не удовлетворяет правилу выравнивания и не заменялся. Завершены 13 допустимых наборов с 78 обучениями и 78 оценками. Во всех 39 наблюдаемых парных сравнениях Hit@32 ниже, чем у случайного контроля. Средние разности отдельных наборов находятся между -0,4343 и -0,0537. Это наблюдение относится только к завершённым наборам. При трёх парах минимальное двустороннее значение точного теста равно 0,25. Результаты разных типов GPU не объединялись в общий эффект.")

    add_heading(doc, "Обучение с бонусом энтропии")
    entropy_method = add_body(doc, "В отдельном опыте три парных seed обучались по 150 шагов с восемью ответами в группе. Опытная ветвь получала бонус 0,01 × H/ln(5) к энтропии выбора операции, контроль обучался тем же кодом без бонуса. Использован групповой нормированный градиент политики без clipping. Его результат нельзя без оговорки переносить на реализации GRPO с clipping.")
    entropy_method.paragraph_format.keep_together = True
    add_body(doc, "Средняя выборочная энтропия группы выросла с 0,7857 до 0,8572 нат. Условная энтропия распределения по 125 программам выросла с 1,4353 до 1,5999 нат. Средний Hit@32 составил 0,4943 в контроле и 0,4887 с бонусом. Средняя парная разность равна -0,0057, 95-процентный t-интервал лежит от -0,0406 до 0,0292, точное двустороннее p равно 0,75. Рост измеренного разнообразия не сопровождался подтверждённым улучшением Hit@32 при данном коэффициенте и бюджете.")

    add_heading(doc, "Текущий вывод")
    add_body(doc, "Структурированный отбор траекторий и проверенный бонус энтропии не улучшили заранее заданный Hit@32. Первые 13 наборов с исключением пар дали отрицательные разности относительно контроля, но эксперимент ещё не завершён. Не начаты 36 обучений в шести допустимых наборах. Вывод о зависимости переноса от числа исключённых пар следует делать после их завершения и отдельного анализа различий между GPU. Латентные представления модели в этих опытах не изучались.")

    doc.core_properties.title = "Отчёт о результатах исследования композиционного обобщения программ"
    doc.core_properties.subject = "Результаты с 22 по 24 сентября 2026 года"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
